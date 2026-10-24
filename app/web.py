"""The web page CSPs and the super admin use.

Run:  python -m app.cli serve        (http://localhost:8000)

Security:
- login with CSP code + password (users.py); forced password change at first login
- session cookie is signed, HttpOnly, SameSite=strict, Secure in production
- 30 minutes without activity logs the user out
- every form carries a CSRF token
- a disabled CSP loses access on their next click (checked on every request)
- pages are never cached by the browser (customer data)

DPDP: customer lists are never saved on the server. A photo / Excel file is read
in memory (reader.py), the rows go to the CSP's phone and stay only there
(templates/csp.html). The server keeps logins, counts per list and the STOP list.
"""
import csv
import io
import json
import logging
import re
import secrets
import time
from datetime import datetime
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import quote, urlencode

from fastapi import FastAPI, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware

from app import config, messages, reader, stats, users
from app.logging_setup import audit, setup_logging
from app.validate import check_rows, group_by_mobile

log = logging.getLogger(__name__)

IDLE_TIMEOUT_SECONDS = 30 * 60
SESSION_MAX_AGE = 8 * 60 * 60

templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))


def u(path):
    """A path inside the app, with the sub-path the app lives under (EKO_BASE_PATH)."""
    return config.BASE_PATH + path
templates.env.globals["retention_days"] = config.RETENTION_DAYS
STATIC_DIR = Path(__file__).parent / "static"
# Eko's official logo is shown in the header once its file is placed here
templates.env.globals["eko_logo"] = (STATIC_DIR / "eko-logo.png").exists()


def _check_secret_key():
    if config.APP_ENV == "production" and (not config.SECRET_KEY or config.SECRET_KEY == "change-me"):
        raise SystemExit("Set a long random SECRET_KEY in .env before running in production")


@asynccontextmanager
async def lifespan(app):
    _check_secret_key()
    setup_logging()
    users.init_db()
    if not any(u["role"] == "admin" for u in users.list_users()):
        log.warning("No super admin yet - run: python -m app.cli create-admin")
    yield


app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
app.add_middleware(
    SessionMiddleware, secret_key=config.SECRET_KEY or secrets.token_urlsafe(32),
    session_cookie="iaa_session", max_age=SESSION_MAX_AGE, same_site="strict",
    path=config.BASE_PATH or "/",
    https_only=config.APP_ENV == "production")


@app.middleware("http")
async def security_headers(request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-store"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["X-Robots-Tag"] = "noindex, nofollow"
    if config.APP_ENV == "production":
        response.headers["Strict-Transport-Security"] = "max-age=31536000"
    return response


# ---------- helpers ----------

class Redirect(Exception):
    def __init__(self, url):
        self.url = url


@app.exception_handler(Redirect)
async def _redirect(request, exc):
    return RedirectResponse(u(exc.url), status_code=303)


class ApiError(Exception):
    def __init__(self, message, status=400):
        self.message, self.status = message, status


@app.exception_handler(ApiError)
async def _api_error(request, exc):
    return JSONResponse({"error": exc.message}, status_code=exc.status)


def _csrf(request):
    if "csrf" not in request.session:
        request.session["csrf"] = secrets.token_urlsafe(24)
    return request.session["csrf"]


def _check_csrf(request, token):
    if not token or not secrets.compare_digest(token, request.session.get("csrf", "")):
        raise Redirect("/login?expired=1")


def _current_user(request, role=None, allow_must_change=False):
    code = request.session.get("user")
    now = time.time()
    if not code or now - request.session.get("seen", 0) > IDLE_TIMEOUT_SECONDS:
        request.session.clear()
        raise Redirect("/login?expired=1" if code else "/login")
    user = users.get_user(code)
    if user is None or not user["active"]:
        request.session.clear()
        raise Redirect("/login")
    request.session["seen"] = now
    if user["must_change"] and not allow_must_change:
        raise Redirect("/password")
    if role and user["role"] != role:
        raise Redirect("/")
    return user


def render(request, name, **ctx):
    ctx.setdefault("user", None)
    ctx.setdefault("base", config.BASE_PATH)
    return templates.TemplateResponse(request, name, {"csrf": _csrf(request), **ctx})


def _fmt_time(iso):
    try:
        dt = datetime.fromisoformat(iso).astimezone()
    except (TypeError, ValueError):
        return ""
    return dt.strftime("%d %b %Y, %I:%M %p")


templates.env.filters["when"] = _fmt_time


# ---------- login ----------

@app.get("/favicon.ico", include_in_schema=False)
def favicon():
    return RedirectResponse(u("/static/icon.svg"), status_code=301)


@app.get("/about", response_class=HTMLResponse)
def about(request: Request):
    """What the agent does - open to everyone, for CSPs and presentations."""
    user = None
    code = request.session.get("user")
    if code and time.time() - request.session.get("seen", 0) <= IDLE_TIMEOUT_SECONDS:
        u = users.get_user(code)
        user = u if u and u["active"] else None
    return render(request, "about.html", user=user)


@app.get("/healthz", include_in_schema=False)
def healthz():
    """For the server's uptime monitor: no login, no data."""
    users.list_users()  # database reachable
    return {"ok": True}


@app.get("/", include_in_schema=False)
def home(request: Request):
    user = _current_user(request)
    return RedirectResponse(u("/admin" if user["role"] == "admin" else "/lists"), status_code=303)


# ---------- first-time setup: create the super admin in the browser ----------

LOCAL_HOSTS = {"127.0.0.1", "::1"}


def _setup_allowed(request):
    """Only before any super admin exists, only from the server's own computer,
    and never in production (there the admin is made with app.cli create-admin) -
    nobody on the internet can claim the admin account."""
    if config.APP_ENV == "production":
        return False
    no_admin = not any(u["role"] == "admin" for u in users.list_users())
    return no_admin and request.client is not None and request.client.host in LOCAL_HOSTS


@app.get("/setup", response_class=HTMLResponse)
def setup_page(request: Request):
    if not _setup_allowed(request):
        return RedirectResponse(u("/login"), status_code=303)
    return render(request, "setup.html")


@app.post("/setup", response_class=HTMLResponse)
def setup_create(request: Request, name: str = Form(""), code: str = Form(""),
                 password: str = Form(""), confirm: str = Form(""), csrf: str = Form("")):
    _check_csrf(request, csrf)
    if not _setup_allowed(request):
        return RedirectResponse(u("/login"), status_code=303)
    try:
        if password != confirm:
            raise users.UserError("The two passwords do not match")
        code = users.create_admin(code, name, password)
    except users.UserError as e:
        return render(request, "setup.html", error=str(e), full_name=name, code=code)
    audit("admin_created", csp_id=code)
    request.session.clear()
    request.session.update({"user": code, "seen": time.time(), "csrf": secrets.token_urlsafe(24)})
    return RedirectResponse(u("/admin"), status_code=303)


def _local_reset_allowed(request):
    """'Forgot admin password' - only on the laptop (never in production), from
    that computer itself. On a server use: python -m app.cli reset-admin-password."""
    return (config.APP_ENV != "production" and request.client is not None
            and request.client.host in LOCAL_HOSTS
            and any(u["role"] == "admin" for u in users.list_users()))


@app.get("/reset-admin", response_class=HTMLResponse)
def reset_admin_page(request: Request):
    if not _local_reset_allowed(request):
        return RedirectResponse(u("/login"), status_code=303)
    admins = [u["csp_code"] for u in users.list_users() if u["role"] == "admin"]
    return render(request, "reset_admin.html", code=admins[0])


@app.post("/reset-admin", response_class=HTMLResponse)
def reset_admin(request: Request, code: str = Form(""), password: str = Form(""),
                confirm: str = Form(""), csrf: str = Form("")):
    _check_csrf(request, csrf)
    if not _local_reset_allowed(request):
        return RedirectResponse(u("/login"), status_code=303)
    try:
        if password != confirm:
            raise users.UserError("The two passwords do not match")
        users.set_admin_password(code, password)
    except users.UserError as e:
        return render(request, "reset_admin.html", code=code, error=str(e))
    code = users.normalise_code(code)
    audit("admin_password_reset", csp_id=code, where="local")
    request.session.clear()
    request.session.update({"user": code, "seen": time.time(), "csrf": secrets.token_urlsafe(24)})
    return RedirectResponse(u("/admin"), status_code=303)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, expired: int = 0):
    if _setup_allowed(request):
        return RedirectResponse(u("/setup"), status_code=303)
    return render(request, "login.html", local_reset=_local_reset_allowed(request),
                  info="Please log in again / कृपया फिर से लॉग इन करें" if expired else None)


@app.post("/login", response_class=HTMLResponse)
def login(request: Request, csp_code: str = Form(""), password: str = Form(""),
          csrf: str = Form("")):
    _check_csrf(request, csrf)
    try:
        user = users.authenticate(csp_code, password)
    except users.UserError as e:
        audit("login_failed", csp_id=users.normalise_code(csp_code)[:32])
        return render(request, "login.html", error=str(e), csp_code=csp_code,
                      local_reset=_local_reset_allowed(request))
    request.session.clear()  # new session id on login
    request.session.update({"user": user["csp_code"], "seen": time.time(),
                            "csrf": secrets.token_urlsafe(24)})
    audit("login", csp_id=user["csp_code"], role=user["role"])
    return RedirectResponse(u("/"), status_code=303)


@app.post("/logout")
def logout(request: Request, csrf: str = Form("")):
    _check_csrf(request, csrf)
    code = request.session.get("user")
    request.session.clear()
    if code:
        audit("logout", csp_id=code)
    return RedirectResponse(u("/login"), status_code=303)


@app.get("/password", response_class=HTMLResponse)
def password_page(request: Request):
    user = _current_user(request, allow_must_change=True)
    return render(request, "password.html", user=user)


@app.post("/password", response_class=HTMLResponse)
def password_change(request: Request, old: str = Form(""), new: str = Form(""),
                    confirm: str = Form(""), csrf: str = Form("")):
    _check_csrf(request, csrf)
    user = _current_user(request, allow_must_change=True)
    try:
        if new != confirm:
            raise users.UserError("The two new passwords do not match")
        users.change_password(user["csp_code"], old, new)
    except users.UserError as e:
        return render(request, "password.html", user=user, error=str(e))
    audit("password_changed", csp_id=user["csp_code"])
    return RedirectResponse(u("/"), status_code=303)


# ---------- CSP: lists (kept only on the CSP's phone) ----------

@app.get("/lists", response_class=HTMLResponse)
def my_lists(request: Request):
    """The CSP's app. The lists themselves are stored in the phone's browser."""
    user = _current_user(request, role="csp")
    return render(request, "csp.html", user=user, followup_days=config.FOLLOWUP_DAYS,
                  max_followups=config.MAX_FOLLOWUPS)


def _api_user(request, csrf=None):
    """Logged-in CSP for the phone's background calls; JSON error instead of a redirect."""
    try:
        user = _current_user(request, role="csp")
    except Redirect:
        raise ApiError("Please log in again / कृपया फिर से लॉग इन करें", 401)
    token = csrf if csrf is not None else request.headers.get("x-csrf-token", "")
    if not token or not secrets.compare_digest(token, request.session.get("csrf", "")):
        raise ApiError("Please reload the page / पेज फिर से खोलें", 401)
    return user


@app.post("/api/read")
async def api_read(request: Request, photo: UploadFile, csrf: str = Form("")):
    """Start reading a photo. Nothing is written to disk."""
    user = _api_user(request, csrf)
    data = await photo.read(reader.MAX_PHOTO_BYTES + 1)
    try:
        token = reader.start(user["csp_code"], data, photo.filename or "photo.jpg")
    except reader.ReadError as e:
        raise ApiError(str(e))
    return {"token": token}


@app.get("/api/read/{token}")
def api_read_status(request: Request, token: str):
    user = _api_user(request)
    return reader.status(token, user["csp_code"])


@app.post("/api/read/{token}/take")
def api_read_take(request: Request, token: str):
    """The phone collects the rows; the server forgets them."""
    user = _api_user(request)
    result = reader.take(token, user["csp_code"])
    if result is None:
        raise ApiError("This reading is not ready or has expired - please upload again")
    rows, info = result
    return {"rows": rows, "info": info}


@app.post("/api/read/{token}/cancel")
def api_read_cancel(request: Request, token: str):
    user = _api_user(request)
    reader.cancel(token, user["csp_code"])
    return {"ok": True}


@app.post("/api/read-sheet")
async def api_read_sheet(request: Request, sheet: UploadFile, csrf: str = Form("")):
    """Excel / CSV list - read in memory, no OCR, nothing kept."""
    user = _api_user(request, csrf)
    data = await sheet.read(reader.MAX_SHEET_BYTES + 1)
    try:
        rows, info = reader.read_sheet(user["csp_code"], data, sheet.filename or "")
    except reader.ReadError as e:
        raise ApiError(str(e))
    return {"rows": rows, "info": info}


def _wa_link(mobile, text):
    return "https://wa.me/91" + mobile + "?" + urlencode({"text": text}, quote_via=quote)


@app.post("/api/check")
async def api_check(request: Request):
    """Check the rows the phone sends and return them with the messages to send.
    Stateless: nothing is stored or logged."""
    user = _api_user(request)
    try:
        body = await request.json()
        raw = []
        for r in body.get("rows", [])[:2000]:
            if isinstance(r, dict):
                row = {k: str(r.get(k) or "")[:120] for k in ("account", "mobile", "name")}
                row.update({k: bool(r.get(k)) for k in ("low_confidence", "unconfirmed")})
                raw.append(row)
    except (ValueError, AttributeError):
        raise ApiError("Bad request")
    rows = check_rows(raw)
    stopped = stats.opted_out(r["mobile"] for r in rows if r["mobile"])
    for r in rows:
        if r["mobile"] in stopped:
            r["stopped"] = True       # said STOP earlier (to any CSP): never messaged
            r["ok"] = False
    groups = group_by_mobile(rows)
    for g in groups:
        g["message"] = messages.render(g, user)
        g["link"] = _wa_link(g["mobile"], g["message"])
        g["reminder_link"] = _wa_link(g["mobile"], messages.render_reminder(g, user))
    counts = {"rows": len(rows), "ok": sum(r["ok"] for r in rows),
              "to_fix": sum(bool(r["problems"]) for r in rows),
              "no_mobile": sum(r["no_mobile"] for r in rows),
              "stopped": sum(bool(r.get("stopped")) for r in rows), "messages": len(groups)}
    return {"rows": rows, "groups": groups, "counts": counts}


@app.post("/api/optout")
async def api_optout(request: Request):
    """Customer said STOP: no CSP will message this number again (hash only)."""
    user = _api_user(request)
    try:
        stats.add_optout((await request.json()).get("mobile"))
    except stats.StatsError as e:
        raise ApiError(str(e))
    except (ValueError, AttributeError):
        raise ApiError("Bad request")
    audit("customer_opted_out", csp_id=user["csp_code"])
    return {"ok": True}


@app.post("/api/stats")
async def api_stats(request: Request):
    """Counts of one list from the phone - numbers only, for the dashboard."""
    user = _api_user(request)
    try:
        body = await request.json()
        counts = body.get("counts") or {}
        stats.record_list(user["csp_code"], body.get("list_id"), counts)
    except stats.StatsError as e:
        raise ApiError(str(e))
    except (ValueError, AttributeError):
        raise ApiError("Bad request")
    if body.get("event") in ("list_saved", "list_deleted", "sending_started"):
        audit(body["event"], csp_id=user["csp_code"],
              **{k: int(counts[k]) for k in ("customers", "to_fix", "messages")
                 if str(counts.get(k, "")).isdigit()})
    return {"ok": True}


# ---------- super admin ----------

def _recent_activity(limit=40):
    path = config.LOG_DIR / "audit.log"
    if not path.exists():
        return []
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()[-limit:]
    out = []
    for line in reversed(lines):
        try:
            out.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return out


ACTIVITY_LABELS = {
    "login": "Logged in", "logout": "Logged out", "login_failed": "Wrong password",
    "password_changed": "Changed password", "list_uploaded": "Uploaded a photo",
    "list_read": "List read (kept on CSP's phone)", "list_saved": "Checked & saved list",
    "list_deleted": "Deleted a list", "sending_started": "Started sending",
    "csps_imported": "Imported CSP sheet",
    "password_reset": "Password reset", "csp_disabled": "CSP disabled",
    "csp_enabled": "CSP enabled",
    "customer_opted_out": "A customer said STOP", "admin_created": "Super admin created",
    "admin_password_reset": "Super admin password reset",
    "customer_blocked": "Number blocked (customer's request)",
    "report_downloaded": "Downloaded CSP report",
}
templates.env.globals["activity_label"] = lambda e: ACTIVITY_LABELS.get(e, e)


def _admin_page(request, user, **ctx):
    csps = [u for u in users.list_users() if u["role"] == "csp"]
    total = stats.totals()
    total["csps"] = sum(u["active"] for u in csps)
    return render(request, "admin.html", user=user, csps=csps, stats=total,
                  progress=stats.per_csp(), earlier=stats.earlier(), activity=_recent_activity(), **ctx)


@app.get("/admin", response_class=HTMLResponse)
def admin(request: Request):
    return _admin_page(request, _current_user(request, role="admin"))


@app.post("/admin/import", response_class=HTMLResponse)
async def admin_import(request: Request, sheet: UploadFile, csrf: str = Form("")):
    _check_csrf(request, csrf)
    user = _current_user(request, role="admin")
    data = await sheet.read(5 * 1024 * 1024 + 1)
    try:
        if len(data) > 5 * 1024 * 1024:
            raise users.UserError("Sheet is too large (max 5 MB)")
        result = users.import_sheet(sheet.filename or "", data)
    except users.UserError as e:
        return _admin_page(request, user, error=str(e))
    audit("csps_imported", csp_id=user["csp_code"], created=len(result["created"]),
          updated=result["updated"], errors=len(result["errors"]))
    return render(request, "passwords.html", user=user, result=result,
                  title="CSP sheet imported")


@app.post("/admin/add-csp", response_class=HTMLResponse)
def admin_add_csp(request: Request, code: str = Form(""), name: str = Form(""),
                  address: str = Form(""), mobile: str = Form(""), csrf: str = Form("")):
    _check_csrf(request, csrf)
    user = _current_user(request, role="admin")
    try:
        result = users.add_csp(code, name, address, mobile)
    except users.UserError as e:
        return _admin_page(request, user, error=str(e))
    audit("csps_imported", csp_id=user["csp_code"], created=1, updated=0, errors=0)
    return render(request, "passwords.html", user=user, result=result, title="CSP added")


@app.post("/admin/erase-customer", response_class=HTMLResponse)
def admin_erase_customer(request: Request, mobile: str = Form(""), csrf: str = Form("")):
    """DPDP: a customer asked not to be contacted. Lists are only on CSP phones;
    the number is blocked here, so no CSP's phone can message it again."""
    _check_csrf(request, csrf)
    user = _current_user(request, role="admin")
    try:
        stats.add_optout(mobile)
    except stats.StatsError as e:
        return _admin_page(request, user, error=str(e))
    audit("customer_blocked", csp_id=user["csp_code"])
    return _admin_page(request, user, notice=(
        "Done. The number is blocked - no CSP can message it again. "
        "(Customer lists are not stored on the server, so there is nothing else to erase here.)"))


@app.get("/guide", response_class=HTMLResponse)
def guide(request: Request):
    """Printable CSP guide (Hindi first) - open to everyone."""
    return render(request, "guide.html")


@app.get("/privacy", response_class=HTMLResponse)
def privacy(request: Request):
    """DPDP notice - open to everyone (customers can be sent this link)."""
    return render(request, "privacy.html", grievance=config.GRIEVANCE_CONTACT)


@app.post("/admin/csp/{code}/reset", response_class=HTMLResponse)
def admin_reset(request: Request, code: str, csrf: str = Form("")):
    _check_csrf(request, csrf)
    user = _current_user(request, role="admin")
    try:
        pw = users.reset_password(code)
    except users.UserError as e:
        return _admin_page(request, user, error=str(e))
    csp = users.get_user(code)
    audit("password_reset", csp_id=csp["csp_code"], by=user["csp_code"])
    return render(request, "passwords.html", user=user, title="Password reset",
                  result={"created": [(csp["csp_code"], csp["name"], pw)], "updated": 0, "errors": []})


@app.post("/admin/csp/{code}/active")
def admin_active(request: Request, code: str, active: int = Form(...), csrf: str = Form("")):
    _check_csrf(request, csrf)
    user = _current_user(request, role="admin")
    try:
        users.set_active(code, bool(active))
    except users.UserError:
        pass
    audit("csp_enabled" if active else "csp_disabled", csp_id=users.normalise_code(code),
          by=user["csp_code"])
    return RedirectResponse(u("/admin#csps"), status_code=303)


@app.get("/admin/report.xlsx")
def admin_report(request: Request):
    """Per-CSP progress as an Excel file. Counts only - no customer data (DPDP)."""
    user = _current_user(request, role="admin")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    progress = stats.per_csp()
    wb = Workbook()
    ws = wb.active
    ws.title = "CSP report"
    ws.append([f"Grahak Sampark - CSP report ({datetime.now().strftime('%d %b %Y, %I:%M %p')})"])
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="0B5CAB")
    ws.append(["Counts only - no customer names or numbers. Repeated lists are counted once."])
    ws["A2"].font = Font(name="Arial", italic=True, color="5D6B7A")
    head = ["CSP code", "CSP name", "Branch", "Status", "Last login", "Different lists", "Customers",
            "Messages to send", "Sent", "Not on WhatsApp", "Said STOP", "Reminders", "Came",
            "Last list"]
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B5CAB")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for u in [u for u in users.list_users() if u["role"] == "csp"]:
        p = progress.get(u["csp_code"], {})
        status = "Disabled" if not u["active"] else ("Not logged in yet" if u["must_change"] else "Active")
        ws.append([
            u["csp_code"], u["name"], u["branch"], status, _fmt_time(u["last_login"]) or "Never",
            p.get("lists", 0) - p.get("repeats", 0), p.get("customers", 0), p.get("messages", 0), p.get("sent", 0),
            p.get("not_on_wa", 0), p.get("stopped", 0), p.get("reminders", 0), p.get("came", 0),
            _fmt_time(p.get("last_list")) if p.get("last_list") else "",
        ])
    for e in stats.earlier():
        ws.append([e["label"], "", "", "", "", e["lists"] - e["repeats"], e["customers"], "", e["sent"], "", "", "",
                   e["came"], ""])
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.font = Font(name="Arial", size=10)
    for col, w in zip("ABCDEFGHIJKLMN", (12, 24, 16, 16, 20, 7, 10, 10, 8, 10, 9, 10, 8, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    audit("report_downloaded", csp_id=user["csp_code"])
    name = f"grahak-sampark-report-{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return Response(buf.getvalue(), headers={"Content-Disposition": f'attachment; filename="{name}"'},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/admin/sheet-template.csv")
def sheet_template(request: Request):
    _current_user(request, role="admin")
    buf = io.StringIO()
    csv.writer(buf).writerow(["CSP Code", "CSP Name", "Shop Address", "Mobile", "Branch"])
    return Response(buf.getvalue(), media_type="text/csv", headers={
        "Content-Disposition": 'attachment; filename="csp-sheet-template.csv"'})
