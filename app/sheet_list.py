"""Read an inoperative-account list from an Excel (.xlsx) or CSV file.

When the CSP has the bank's list as a file, no OCR is needed: the cells are
read directly, so every digit is exact. The same rules as for photos apply:

- columns are found by their titles (same spellings as ocr_table.HEADER_KEYWORDS);
- only INOPERATIVE customers are kept when there is an Operative/Inoperative column;
- only account, mobile and name are kept (DPDP minimisation) - the file itself
  is read in memory and never stored.
"""
import csv
import io
import re
from pathlib import Path

from app.ocr_table import _header_matches

ALLOWED = {".xlsx", ".csv"}
MAX_BYTES = 5 * 1024 * 1024
HEADER_SEARCH_ROWS = 15  # the title row may come after a few lines of bank heading


class SheetError(Exception):
    pass


def _cell_text(v):
    """Excel stores long numbers as int/float (and may show 3.37E+10); turn them
    back into plain digits."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows_from_file(filename, data):
    ext = Path(filename).suffix.lower()
    if ext == ".xls":
        raise SheetError("Old Excel format (.xls). Please open it and 'Save as' .xlsx, then upload")
    if ext not in ALLOWED:
        raise SheetError("Please upload an Excel (.xlsx) or CSV file")
    if not data:
        raise SheetError("The file is empty")
    if len(data) > MAX_BYTES:
        raise SheetError("File is too large (max 5 MB)")
    if ext == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise SheetError("Could not open this Excel file")
    try:
        return [[_cell_text(v) for v in r] for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def _find_header(rows):
    for i, r in enumerate(rows[:HEADER_SEARCH_ROWS]):
        cols = {}
        for j, cell in enumerate(r):
            for field, _ in _header_matches(cell or "", 0, 1):
                cols.setdefault(field, j)
                break  # one field per cell
        if "account" in cols and "mobile" in cols:
            return i, cols
    return None, None


def read_customer_sheet(filename, data):
    """-> (rows, info) in the same shape as ocr_table.extract_rows."""
    rows = _rows_from_file(filename, data)
    h, cols = _find_header(rows)
    if h is None:
        raise SheetError("Could not find the Account and Mobile column titles in the file")
    out, skipped = [], 0

    def cell(r, field):
        j = cols.get(field)
        return r[j] if j is not None and j < len(r) else ""

    for r in rows[h + 1:]:
        if not any(r):
            continue
        account = re.sub(r"\D", "", cell(r, "account"))
        mobile = re.sub(r"\D", "", cell(r, "mobile"))   # "9876543210BB" -> 9876543210
        if not account and not mobile:
            continue  # totals line, notes, blank row
        status = cell(r, "_status").upper()
        if status.startswith("OPER"):
            skipped += 1  # operative account - not a customer to contact, not kept
            continue
        out.append({"account": account, "mobile": mobile, "name": cell(r, "name"),
                    "low_confidence": False})
    if not out:
        raise SheetError("No inoperative customers found in the file"
                         + (f" ({skipped} operative accounts skipped)" if skipped else ""))
    return out, {"columns": sorted(f for f in cols if not f.startswith("_")),
                 "skipped_operative": skipped}
