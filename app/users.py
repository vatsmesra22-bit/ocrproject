"""CSP and super-admin accounts, stored in SQLite (DATA_DIR/app.db).

- Login ID is the CSP code. The super admin is created once from the command
  line (python -m app.cli create-admin), never through the web.
- Passwords are stored only as scrypt hashes. Temporary passwords are shown to
  the super admin once and must be changed at first login.
- Wrong passwords lock the account for a while (LOGIN_* settings; 0 = no lock).
- CSP accounts are created/updated by importing the super admin's sheet.
Table definitions live in db.py.
"""
import csv
import hashlib
import hmac
import io
import re
import secrets
from datetime import datetime, timedelta, timezone

from app import config, db

MIN_PASSWORD_LEN = 8

class UserError(Exception):
    pass


def _now():
    return datetime.now(timezone.utc)


def _iso(dt):
    return dt.isoformat(timespec="seconds")


def init_db():
    db.init_db()


def normalise_code(code):
    return re.sub(r"\s+", "", str(code or "")).upper()


# ---------- passwords ----------

def hash_password(password):
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode(), salt=salt, n=2 ** 14, r=8, p=1)
    return f"scrypt$16384$8$1${salt.hex()}${digest.hex()}"


def verify_password(password, stored):
    try:
        _, n, r, p, salt, digest = stored.split("$")
        check = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt),
                               n=int(n), r=int(r), p=int(p))
    except (ValueError, TypeError):
        return False
    return hmac.compare_digest(check.hex(), digest)


def check_password_strength(password):
    if len(password) < MIN_PASSWORD_LEN:
        raise UserError(f"Password must be at least {MIN_PASSWORD_LEN} characters")
    if not re.search(r"[A-Za-z]", password) or not re.search(r"\d", password):
        raise UserError("Password must contain both letters and numbers")


def temp_password():
    # no 0/O/1/l/I so it can be read out over the phone
    alphabet = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    while True:
        pw = "".join(secrets.choice(alphabet) for _ in range(10))
        if re.search(r"\d", pw) and re.search(r"[A-Za-z]", pw):
            return pw


# ---------- accounts ----------

def get_user(csp_code):
    with db.connect() as conn:
        row = conn.execute("SELECT * FROM users WHERE csp_code = ?",
                           (normalise_code(csp_code),)).fetchone()
    return dict(row) if row else None


def list_users():
    with db.connect() as conn:
        rows = conn.execute("SELECT * FROM users ORDER BY role, csp_code").fetchall()
    return [dict(r) for r in rows]


def create_admin(code, name, password):
    code = normalise_code(code)
    if not code:
        raise UserError("Admin login ID is required")
    check_password_strength(password)
    now = _iso(_now())
    with db.connect() as conn:
        if conn.execute("SELECT 1 FROM users WHERE csp_code = ?", (code,)).fetchone():
            raise UserError(f"Login ID {code} already exists")
        conn.execute(
            "INSERT INTO users (csp_code, role, name, pw_hash, must_change, created_at, updated_at)"
            " VALUES (?, 'admin', ?, ?, 0, ?, ?)", (code, name or "Super Admin",
                                                     hash_password(password), now, now))
    return code


def authenticate(csp_code, password):
    """Returns the user dict, or raises UserError with a message safe to show."""
    user = get_user(csp_code)
    generic = UserError("Wrong login ID or password")
    if user is None:
        hash_password(password)  # same time as a real check: don't reveal which IDs exist
        raise generic
    now = _now()
    if user["locked_until"] and datetime.fromisoformat(user["locked_until"]) > now:
        raise UserError(f"Too many wrong attempts. Try again after {config.LOGIN_LOCKOUT_MINUTES} minutes.")
    if not verify_password(password, user["pw_hash"]):
        failed = user["failed_logins"] + 1
        lock_on = config.LOGIN_LOCKOUT_MINUTES > 0 and failed >= config.LOGIN_MAX_FAILED
        locked = _iso(now + timedelta(minutes=config.LOGIN_LOCKOUT_MINUTES)) if lock_on else None
        with db.connect() as conn:
            conn.execute("UPDATE users SET failed_logins = ?, locked_until = ? WHERE csp_code = ?",
                         (0 if locked else failed, locked, user["csp_code"]))
        raise generic
    if not user["active"]:
        raise UserError("This account is disabled. Please contact Eko.")
    with db.connect() as conn:
        conn.execute("UPDATE users SET failed_logins = 0, locked_until = NULL, last_login = ?"
                     " WHERE csp_code = ?", (_iso(now), user["csp_code"]))
    return get_user(user["csp_code"])


def change_password(csp_code, old, new):
    user = get_user(csp_code)
    if user is None or not verify_password(old, user["pw_hash"]):
        raise UserError("Current password is wrong")
    if old == new:
        raise UserError("New password must be different")
    check_password_strength(new)
    with db.connect() as conn:
        conn.execute("UPDATE users SET pw_hash = ?, must_change = 0, updated_at = ?"
                     " WHERE csp_code = ?", (hash_password(new), _iso(_now()), user["csp_code"]))


def set_admin_password(csp_code, new):
    """Laptop-only 'forgot admin password' page (web.py) and the server CLI."""
    user = get_user(csp_code)
    if user is None or user["role"] != "admin":
        raise UserError("No super admin with this login ID")
    check_password_strength(new)
    with db.connect() as conn:
        conn.execute("UPDATE users SET pw_hash = ?, must_change = 0, failed_logins = 0,"
                     " locked_until = NULL, updated_at = ? WHERE csp_code = ?",
                     (hash_password(new), _iso(_now()), user["csp_code"]))


def reset_password(csp_code):
    """Super admin action: new temporary password, returned once."""
    user = get_user(csp_code)
    if user is None or user["role"] != "csp":
        raise UserError("CSP not found")
    pw = temp_password()
    with db.connect() as conn:
        conn.execute("UPDATE users SET pw_hash = ?, must_change = 1, failed_logins = 0,"
                     " locked_until = NULL, updated_at = ? WHERE csp_code = ?",
                     (hash_password(pw), _iso(_now()), user["csp_code"]))
    return pw


def set_active(csp_code, active):
    user = get_user(csp_code)
    if user is None or user["role"] != "csp":
        raise UserError("CSP not found")
    with db.connect() as conn:
        conn.execute("UPDATE users SET active = ?, updated_at = ? WHERE csp_code = ?",
                     (1 if active else 0, _iso(_now()), user["csp_code"]))


def add_csp(code, name, address, mobile="", branch=""):
    """Super admin adds one CSP by hand (no sheet). Same rules as the import."""
    buf = io.StringIO()
    csv.writer(buf).writerows([["CSP Code", "CSP Name", "Shop Address", "Mobile", "Branch"],
                               [code, name, address, mobile, branch]])
    result = import_sheet("one.csv", buf.getvalue().encode())
    if result["errors"]:
        raise UserError(result["errors"][0].split(": ", 1)[-1])
    if result["updated"]:
        raise UserError(f"{normalise_code(code)} already exists - use Reset password if needed")
    return result


# ---------- sheet import ----------

SHEET_COLUMNS = {
    "csp_code": ["CSPCODE", "CSPID", "AGENTCODE", "AGENTID", "CODE", "BCAGENTID"],
    "name": ["CSPNAME", "AGENTNAME", "OUTLETNAME", "SHOPNAME", "NAME"],
    "address": ["SHOPADDRESS", "ADDRESS", "VILLAGE", "LOCATION", "ADDR"],
    "mobile": ["MOBILENO", "MOBILENUMBER", "MOBILE", "PHONE", "CONTACTNO", "CONTACT"],
    "branch": ["BRANCHNAME", "BRANCH", "REGION"],
}
REQUIRED = ("csp_code", "name", "address")


def _read_sheet(filename, data):
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if v is None else str(v) for v in r] for r in wb.active.iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    raise UserError("Please upload the sheet as .xlsx or .csv")


def _map_header(header):
    norm = [re.sub(r"[^A-Z0-9]", "", h.upper()) for h in header]
    mapping = {}
    for field, names in SHEET_COLUMNS.items():
        for kw in names:  # most specific spelling first
            for i, h in enumerate(norm):
                if h == kw and i not in mapping.values():
                    mapping[field] = i
                    break
            if field in mapping:
                break
    return mapping


def import_sheet(filename, data):
    """Create new CSPs and update existing ones. Returns
    {"created": [(code, name, temp_password)], "updated": n, "errors": [str]}."""
    rows = [r for r in _read_sheet(filename, data) if any(str(c).strip() for c in r)]
    if not rows:
        raise UserError("The sheet is empty")
    mapping = _map_header(rows[0])
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        raise UserError("Sheet is missing column(s): " + ", ".join(
            m.replace("_", " ").upper() for m in missing))

    created, updated, errors, seen = [], 0, [], set()
    now = _iso(_now())
    with db.connect() as conn:
        for line_no, r in enumerate(rows[1:], start=2):
            def cell(field):
                i = mapping.get(field)
                return str(r[i]).strip() if i is not None and i < len(r) else ""
            code = normalise_code(cell("csp_code"))
            mobile = re.sub(r"\D", "", cell("mobile"))
            if mobile.startswith("91") and len(mobile) == 12:
                mobile = mobile[2:]
            if not code:
                errors.append(f"Row {line_no}: CSP code is empty")
                continue
            if code in seen:
                errors.append(f"Row {line_no}: {code} appears twice in the sheet")
                continue
            seen.add(code)
            if not cell("name") or not cell("address"):
                errors.append(f"Row {line_no}: {code} needs both name and address")
                continue
            existing = conn.execute("SELECT role FROM users WHERE csp_code = ?", (code,)).fetchone()
            if existing and existing["role"] == "admin":
                errors.append(f"Row {line_no}: {code} is the admin login - skipped")
                continue
            values = (cell("name"), cell("address"), mobile, cell("branch"), now, code)
            if existing:
                conn.execute("UPDATE users SET name = ?, address = ?, mobile = ?, branch = ?,"
                             " updated_at = ? WHERE csp_code = ?", values)
                updated += 1
            else:
                pw = temp_password()
                conn.execute(
                    "INSERT INTO users (name, address, mobile, branch, updated_at, csp_code,"
                    " role, pw_hash, must_change, created_at) VALUES (?, ?, ?, ?, ?, ?, 'csp', ?, 1, ?)",
                    values + (hash_password(pw), now))
                created.append((code, cell("name"), pw))
    return {"created": created, "updated": updated, "errors": errors}
