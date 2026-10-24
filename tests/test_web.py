"""Web flow tests with FastAPI's test client. OCR is stubbed out.

Customer lists are never saved on the server: the photo is read in memory, the
rows go to the CSP's phone, and the server keeps only counts and the STOP list.
"""
import io
import re
import time

import pytest
from fastapi.testclient import TestClient

from app import config, reader, stats, users, web

ROWS = [{"account": "11111111111", "mobile": "9876543210", "name": "SITA DEVI", "low_confidence": False},
        {"account": "2222222222", "mobile": "9876543211", "name": "RAM LAL", "low_confidence": False}]


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(config, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(reader, "extract_rows", lambda data: (ROWS, {"skipped_operative": 1}))
    users.init_db()
    users.create_admin("BOSS", "Admin", "adminpass1")
    created = users.import_sheet("x.csv", b"CSP Code,CSP Name,Address\nA1,Kendra A,Village A\n"
                                          b"B2,Kendra B,Village B\n")["created"]
    pws = {code: pw for code, _, pw in created}
    with TestClient(web.app) as client:
        yield client, pws


def csrf(client, url="/login"):
    return re.search(r'name="csrf" value="([^"]+)"', client.get(url).text).group(1)


def login(client, code, pw):
    return client.post("/login", data={"csp_code": code, "password": pw, "csrf": csrf(client)})


def login_csp(client, code, pw):
    login(client, code, pw)
    client.post("/password", data={"old": pw, "new": "newpass123", "confirm": "newpass123",
                                   "csrf": csrf(client, "/password")})
    return {"X-CSRF-Token": csrf(client, "/lists")}


def read_photo(client, h):
    token = client.post("/api/read", data={"csrf": h["X-CSRF-Token"]},
                        files={"photo": ("list.jpg", b"jpeg-bytes", "image/jpeg")}).json()["token"]
    for _ in range(100):
        if client.get(f"/api/read/{token}", headers=h).json()["stage"] == "done":
            break
        time.sleep(0.05)
    return token


def test_pages_need_login(env):
    client, _ = env
    for url in ("/", "/lists", "/admin"):
        r = client.get(url, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"].startswith("/login")
    r = client.post("/api/check", json={"rows": ROWS})
    assert r.status_code == 401 and "log in" in r.json()["error"]


def test_wrong_password_and_missing_csrf(env):
    client, pws = env
    assert "Wrong login ID or password" in login(client, "A1", "nope").text
    r = client.post("/login", data={"csp_code": "A1", "password": pws["A1"]}, follow_redirects=False)
    assert r.headers["location"].startswith("/login")


def test_api_needs_csrf_token(env):
    client, pws = env
    login_csp(client, "A1", pws["A1"])
    assert client.post("/api/check", json={"rows": ROWS}).status_code == 401
    assert client.post("/api/check", json={"rows": ROWS}, headers={"X-CSRF-Token": "x"}).status_code == 401


def test_first_login_forces_password_change(env):
    client, pws = env
    r = login(client, "A1", pws["A1"])
    assert "Set your own password" in r.text
    assert client.get("/lists", follow_redirects=False).headers["location"] == "/password"


def test_photo_read_in_memory_and_nothing_kept_on_server(env):
    client, pws = env
    h = login_csp(client, "A1", pws["A1"])
    token = read_photo(client, h)
    d = client.post(f"/api/read/{token}/take", headers=h).json()
    assert [r["account"] for r in d["rows"]] == ["11111111111", "2222222222"]
    assert d["info"]["skipped_operative"] == 1
    # collected once - then gone from the server's memory
    assert client.post(f"/api/read/{token}/take", headers=h).status_code == 400
    assert client.get(f"/api/read/{token}", headers=h).json()["stage"] == "gone"
    # no photo, list or customer number anywhere on disk
    files = [p for p in config.DATA_DIR.rglob("*") if p.is_file()]
    assert [p.name for p in files] == ["app.db"]
    blob = config.DATA_DIR.joinpath("app.db").read_bytes()
    logs = "".join(p.read_text(encoding="utf-8") for p in config.LOG_DIR.glob("*.log"))
    for secret in ("9876543210", "SITA", "11111111111"):
        assert secret.encode() not in blob and secret not in logs


def test_another_csp_cannot_collect_my_reading(env):
    client, pws = env
    h = login_csp(client, "A1", pws["A1"])
    token = read_photo(client, h)
    client.post("/logout", data={"csrf": csrf(client, "/lists")})
    h2 = login_csp(client, "B2", pws["B2"])
    assert client.post(f"/api/read/{token}/take", headers=h2).status_code == 400
    assert client.get(f"/api/read/{token}", headers=h2).json()["stage"] == "gone"


def test_unreadable_photo_reports_error(env, monkeypatch):
    client, pws = env
    monkeypatch.setattr(reader, "extract_rows", lambda data: ([], {"error": "Photo is not clear"}))
    h = login_csp(client, "A1", pws["A1"])
    token = client.post("/api/read", data={"csrf": h["X-CSRF-Token"]},
                        files={"photo": ("list.jpg", b"x", "image/jpeg")}).json()["token"]
    for _ in range(100):
        s = client.get(f"/api/read/{token}", headers=h).json()
        if s["stage"] == "failed":
            break
        time.sleep(0.05)
    assert s["stage"] == "failed" and "not clear" in s["error"]
    bad = client.post("/api/read", data={"csrf": h["X-CSRF-Token"]},
                      files={"photo": ("list.pdf", b"x", "application/pdf")})
    assert bad.status_code == 400 and "JPG or PNG" in bad.json()["error"]


def test_check_flags_problems_and_builds_whatsapp_links(env):
    client, pws = env
    h = login_csp(client, "A1", pws["A1"])
    d = client.post("/api/check", json={"rows": ROWS + [
        {"account": "33333333333", "mobile": "", "name": "NO PHONE"}]}, headers=h).json()
    assert d["counts"] == {"rows": 3, "ok": 1, "to_fix": 1, "no_mobile": 1, "stopped": 0, "messages": 1}
    assert "account should be 11 digits" in d["rows"][1]["problems"]
    g = d["groups"][0]
    assert g["link"].startswith("https://wa.me/919876543210?text=") and "Kendra%20A" in g["link"]
    assert "XXXX1111" in g["message"] and g["reminder_link"].startswith("https://wa.me/91")


def test_stop_blocks_the_number_for_every_csp(env):
    client, pws = env
    h = login_csp(client, "A1", pws["A1"])
    assert client.post("/api/optout", json={"mobile": "9876543210"}, headers=h).json() == {"ok": True}
    client.post("/logout", data={"csrf": csrf(client, "/lists")})
    h2 = login_csp(client, "B2", pws["B2"])
    d = client.post("/api/check", json={"rows": ROWS[:1]}, headers=h2).json()
    assert d["rows"][0]["stopped"] and d["groups"] == [] and d["counts"]["stopped"] == 1


def test_csp_uploads_excel_file_read_in_memory(env):
    client, pws = env
    h = login_csp(client, "A1", pws["A1"])
    assert "Upload Excel file" in client.get("/lists").text
    csv_data = (b"A/C No,Name,Operative/Inoperative,Mobile No\n"
                b"11111111111,SITA DEVI,INOPERATIVE,9876543210\n"
                b"22222222222,RAM LAL,OPERATIVE,9876543211\n")
    d = client.post("/api/read-sheet", data={"csrf": h["X-CSRF-Token"]},
                    files={"sheet": ("list.csv", csv_data, "text/csv")}).json()
    assert len(d["rows"]) == 1 and d["info"] == {"skipped_operative": 1, "source": "sheet"}
    assert [p.name for p in config.DATA_DIR.rglob("*") if p.is_file()] == ["app.db"]
    bad = client.post("/api/read-sheet", data={"csrf": h["X-CSRF-Token"]},
                      files={"sheet": ("list.xls", b"x", "application/vnd.ms-excel")})
    assert bad.status_code == 400 and "Save as" in bad.json()["error"]


def test_phone_counts_feed_dashboard_and_report(env):
    from openpyxl import load_workbook
    client, pws = env
    h = login_csp(client, "A1", pws["A1"])
    post = lambda lid, **c: client.post("/api/stats", json={"list_id": lid, "event": "list_saved", "counts": c},
                                        headers=h)
    assert post("aaaaaaaaaaaa", customers=10, messages=8, sent=3, came=1, stopped=1).json() == {"ok": True}
    post("bbbbbbbbbbbb", customers=10, messages=8, sent=5, repeat=1)        # same list again
    post("cccccccccccc", customers=4, deleted=1)                             # deleted on the phone
    assert post("../etc", customers=1).status_code == 400
    client.post("/logout", data={"csrf": csrf(client, "/lists")})
    login(client, "BOSS", "adminpass1")
    page = client.get("/admin").text
    assert "1</b><span>Different lists tested" in page and "2 uploads in all - 1 were the same list again" in page
    assert "10</b><span>Customers on lists" in page and "8</b><span>WhatsApp sent" in page
    ws = load_workbook(io.BytesIO(client.get("/admin/report.xlsx").content)).active
    rows = {row[0]: row for row in ws.iter_rows(min_row=5, values_only=True)}
    assert rows["A1"][5:13] == (1, 10, 8, 8, 0, 1, 0, 1)
    assert rows["B2"][5:13] == (0, 0, 0, 0, 0, 0, 0, 0)


def test_admin_blocks_a_number(env):
    client, pws = env
    login(client, "BOSS", "adminpass1")
    r = client.post("/admin/erase-customer", data={"mobile": "+91 98765 43210", "csrf": csrf(client, "/admin")})
    assert "number is blocked" in r.text
    assert stats.opted_out(["9876543210"]) == {"9876543210"}
    r = client.post("/admin/erase-customer", data={"mobile": "123", "csrf": csrf(client, "/admin")})
    assert "10-digit" in r.text


def test_csp_page_says_lists_stay_on_the_phone(env):
    client, pws = env
    login_csp(client, "A1", pws["A1"])
    page = client.get("/lists").text
    assert "only on this phone" in page and 'href="/guide"' in page and "localStorage" in page


def test_admin_import_reset_disable(env):
    client, pws = env
    login(client, "BOSS", "adminpass1")
    page = client.get("/admin").text
    assert "Super admin" in page and "A1" in page
    r = client.post("/admin/import", data={"csrf": csrf(client, "/admin")},
                    files={"sheet": ("s.csv", b"CSP Code,CSP Name,Address\nC3,Kendra C,Village C\n", "text/csv")})
    assert "shown <u>only once</u>" in r.text and "C3" in r.text
    r = client.post("/admin/csp/A1/reset", data={"csrf": csrf(client, "/admin")})
    assert "Password reset" in r.text
    client.post("/admin/csp/B2/active", data={"active": "0", "csrf": csrf(client, "/admin")})
    assert not users.get_user("B2")["active"]


def test_disabled_csp_is_logged_out_on_next_click(env):
    client, pws = env
    login_csp(client, "A1", pws["A1"])
    users.set_active("A1", False)
    assert client.get("/lists", follow_redirects=False).headers["location"] == "/login"


def test_idle_session_expires(env, monkeypatch):
    client, pws = env
    login_csp(client, "A1", pws["A1"])
    real = time.time
    monkeypatch.setattr(web.time, "time", lambda: real() + web.IDLE_TIMEOUT_SECONDS + 5)
    assert client.get("/lists", follow_redirects=False).headers["location"].startswith("/login?expired")


def test_security_headers(env):
    client, _ = env
    r = client.get("/login")
    assert r.headers["cache-control"] == "no-store" and r.headers["x-frame-options"] == "DENY"


def test_browser_setup_only_before_admin_exists(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(config, "LOG_DIR", tmp_path / "logs")
    users.init_db()
    monkeypatch.setattr(web, "LOCAL_HOSTS", {"testclient"})  # the test client counts as "this computer"
    with TestClient(web.app) as client:
        assert client.get("/login", follow_redirects=False).headers["location"] == "/setup"
        tok = csrf(client, "/setup")
        r = client.post("/setup", data={"name": "Admin", "code": "boss", "password": "lettersonly",
                                        "confirm": "lettersonly", "csrf": tok})
        assert "letters and numbers" in r.text
        r = client.post("/setup", data={"name": "Admin", "code": "boss", "password": "goodpass12",
                                        "confirm": "goodpass12", "csrf": tok})
        assert "Super admin" in r.text and users.get_user("BOSS")["role"] == "admin"
        assert client.get("/setup", follow_redirects=False).headers["location"] == "/login"


def test_browser_setup_refused_from_other_computers(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(config, "LOG_DIR", tmp_path / "logs")
    users.init_db()
    with TestClient(web.app) as client:  # "testclient" is not a local address
        assert client.get("/setup", follow_redirects=False).headers["location"] == "/login"
        assert "Welcome" in client.get("/login").text


def test_setup_page_never_in_production(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(config, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(config, "APP_ENV", "production")
    monkeypatch.setattr(config, "SECRET_KEY", "x" * 40)
    monkeypatch.setattr(web, "LOCAL_HOSTS", {"testclient"})
    users.init_db()
    with TestClient(web.app) as client:
        assert client.get("/setup", follow_redirects=False).headers["location"] == "/login"
        assert client.get("/healthz").json() == {"ok": True}


def test_local_admin_password_reset(env, monkeypatch):
    client, _ = env
    monkeypatch.setattr(web, "LOCAL_HOSTS", {"testclient"})
    assert "password bhool gaye" in client.get("/login").text
    tok = csrf(client, "/reset-admin")
    r = client.post("/reset-admin", data={"code": "BOSS", "password": "brandnew77",
                                         "confirm": "brandnew77", "csrf": tok})
    assert "Super admin" in r.text
    assert users.authenticate("BOSS", "brandnew77")


def test_admin_reset_page_closed_from_other_computers(env):
    client, _ = env  # "testclient" is not a local address here
    assert "password bhool gaye" not in client.get("/login").text
    assert client.get("/reset-admin", follow_redirects=False).headers["location"] == "/login"


def test_admin_adds_one_csp_without_a_sheet(env):
    client, _ = env
    login(client, "BOSS", "adminpass1")
    r = client.post("/admin/add-csp", data={"code": "c9", "name": "Kendra Nine", "address": "Village 9",
                                           "csrf": csrf(client, "/admin")})
    assert "C9" in r.text and "only once" in r.text
    assert users.get_user("C9")["must_change"] == 1
    r = client.post("/admin/add-csp", data={"code": "C9", "name": "Again", "address": "X",
                                           "csrf": csrf(client, "/admin")})
    assert "already exists" in r.text


def test_about_page_is_public(env):
    client, _ = env
    r = client.get("/about")
    assert r.status_code == 200 and "4 simple steps" in r.text and "Log in" in r.text
    assert 'href="/about"' in client.get("/login").text


def test_privacy_page_and_grievance_line(env, monkeypatch):
    from app import messages
    client, _ = env
    assert "Grievance officer contact to be added" in client.get("/privacy").text
    monkeypatch.setattr(config, "GRIEVANCE_CONTACT", "grievance@example.org")
    assert "grievance@example.org" in client.get("/privacy").text
    text = messages.render({"greeting": "Sita", "accounts_last4": ["1111"]}, {"name": "K", "address": "V"})
    assert text.endswith("Questions / complaints: grievance@example.org")
    assert 'href="/privacy"' in client.get("/login").text


def test_csp_guide_is_public(env):
    client, _ = env
    r = client.get("/guide")
    assert r.status_code == 200 and "CSP गाइड" in r.text


def test_earlier_work_still_counted_on_dashboard(env):
    client, _ = env
    stats.add_earlier("Earlier tests (before 23 Oct 2026)", lists=21, repeats=12, unread=3, customers=341)
    login(client, "BOSS", "adminpass1")
    page = client.get("/admin").text
    assert "9</b><span>Different lists tested" in page and "341</b><span>Customers on lists" in page
    assert "3 upload(s) could not be read" in page
    assert "Earlier tests (before 23 Oct 2026)" in page


def test_app_under_a_sub_path(tmp_path, monkeypatch):
    """EKO_BASE_PATH=/ocrproject: every link, redirect and page call stays inside it."""
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(config, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(config, "BASE_PATH", "/ocrproject")
    users.init_db()
    users.create_admin("BOSS", "Admin", "adminpass1")
    pw = users.import_sheet("x.csv", b"CSP Code,CSP Name,Address\nA1,Kendra A,Village A\n")["created"][0][2]
    with TestClient(web.app, root_path="/ocrproject") as client:
        r = client.get("/ocrproject/lists", follow_redirects=False)
        assert r.headers["location"] == "/ocrproject/login"
        page = client.get("/ocrproject/login").text
        assert 'action="/ocrproject/login"' in page and 'href="/ocrproject/static/icon.svg"' in page
        assert not re.search(r'(href|action|src)="/(?!ocrproject|/)', page)
        tok = re.search(r'name="csrf" value="([^"]+)"', page).group(1)
        r = client.post("/ocrproject/login", data={"csp_code": "A1", "password": pw, "csrf": tok},
                        follow_redirects=False)
        assert r.headers["location"] == "/ocrproject/"
        assert client.get("/ocrproject/", follow_redirects=False).headers["location"] == "/ocrproject/password"
        tok = re.search(r'name="csrf" value="([^"]+)"', client.get("/ocrproject/password").text).group(1)
        client.post("/ocrproject/password", data={"old": pw, "new": "newpass123", "confirm": "newpass123", "csrf": tok})
        page = client.get("/ocrproject/lists").text
        assert 'var BASE = "/ocrproject"' in page and "fetch(BASE + url" in page
        assert not re.search(r'(href|action|src)="/(?!ocrproject|/)', page)
        tok = re.search(r'name="csrf" value="([^"]+)"', page).group(1)
        d = client.post("/ocrproject/api/check", json={"rows": ROWS}, headers={"X-CSRF-Token": tok}).json()
        assert d["counts"]["rows"] == 2
        r = client.post("/ocrproject/logout", data={"csrf": tok}, follow_redirects=False)
        assert r.headers["location"] == "/ocrproject/login"
